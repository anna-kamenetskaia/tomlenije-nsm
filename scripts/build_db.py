import sqlite3, json, re, csv, pandas as pd
DB='out/kmiha_project.db'
import os; os.path.exists(DB) and os.remove(DB)
con=sqlite3.connect(DB); cur=con.cursor()
cur.executescript(open('schema.sql').read())

# 1. corpus
df=pd.read_pickle('df.pkl')
corp=df[['example_id','source','kwic','left','right','full_text','dom']].rename(columns={'kwic':'form','left':'left_ctx','right':'right_ctx','dom':'domain_auto'})
corp['domain_auto']=corp['domain_auto'].replace('none',None)
corp.to_sql('corpus',con,if_exists='append',index=False)

# 2. contexts
c=pd.read_csv('contexts_24.csv',encoding='utf-8-sig')
c.rename(columns={'n':'position','domain_hint':'domain','text':'text_in_prompt'})[['position','example_id','domain','text_in_prompt']].to_sql('prompt_contexts',con,if_exists='append',index=False)

# 3. primes
def cats(f,start,end):
    s=open(f).read(); a=s.index(start); b=s.index('\n\n',s.index(end))
    return [(l.split(':',1)[0].strip(),[p.strip() for p in re.split(r',\s*(?![^()]*\))',l.split(':',1)[1])]) for l in s[a:b].strip().split('\n')]
ru=cats('prompt_ru.txt','Субстантивы:','Сходство:'); en=cats('prompt_en.txt','Substantives:','Similarity:'); he=cats('prompt_he.txt','שמות עצם:','דמיון:')
n=0
for (cr,a),(ce,b),(ch,d) in zip(ru,en,he):
    for x,y,z in zip(a,b,d):
        n+=1; cur.execute('insert into primes values (?,?,?,?,?,?)',(n,ce,cr,y,x,z))

# 4. prompts
for lang in ['ru','en','he']:
    cur.execute('insert into prompts values (?,?)',(lang,open(f'prompt_{lang}.txt').read()))

# 5. models (вкл. отсеянные)
models=[('Nemotron3','nvidia/nemotron-3-super-120b-a12b:free','NVIDIA','США','использована',None),
('Dots3','dots-studio/dots-3-note-preview:free','Xiaohongshu (dots.llm)','Китай','использована',None),
('Gemma4','google/gemma-4-31b-it:free','Google','США','отсеяна','429: бесплатный пул Google перегружен'),
('Inkling','thinkingmachines/inkling:free','Thinking Machines','США','отсеяна','403: доступна только в агентных приложениях'),
('Xing4','xingchen-agi/xing4.0-29b-a4b:free','xingchen-agi',None,'отсеяна','~10 мин на ответ, искажение кириллицы в рассуждениях')]
cur.executemany('insert into models values (?,?,?,?,?,?)',models)

# 6. runs
res=json.load(open('results.json'))
key={(r['model'],r['lang'],int(r['iter'])):r['code'] for r in csv.DictReader(open('out/key_DO_NOT_OPEN.csv',encoding='utf-8-sig'))}
GARBLED={('Dots3','he',1),('Dots3','he',3),('Dots3','he',4),('Dots3','he',5),('Nemotron3','ru',3)}
LOOP={('Dots3','ru',2),('Dots3','ru',4)}
for r in res:
    k=(r['model'],r['lang'],r['iter']); p=r['parsed'] or {}
    status='отказ' if not r['raw'] else ('не тот язык / смесь' if k in GARBLED else ('зацикливание' if k in LOOP else 'ok'))
    ex=p.get('explication',[])
    cur.execute('insert into runs (code,model,lang,iter,temperature,max_tokens,status,spheres_json,explication,n_lines,raw_response) values (?,?,?,?,?,?,?,?,?,?,?)',
        (key[k],r['model'],r['lang'],r['iter'],r['temperature'],16000,status,json.dumps(p.get('spheres',[]),ensure_ascii=False) if p else None,
         '\n'.join(ex) if ex else None,len(ex),r['raw']))

# 7. scores (К1 из таблицы; К2–К5 — пусто до экспертной оценки)
from openpyxl import load_workbook
ws=load_workbook('out/scoring.xlsx',data_only=True)['Оценка']
for row in ws.iter_rows(min_row=3,values_only=True):
    code,lang,_,_,k1,k2,k3,k4,k5,_,comment,why=row[:12]
    cur.execute('insert into scores (code,k1,k1_reason,k2,k3,k4,k5,comment) values (?,?,?,?,?,?,?,?)',(code,k1,why,k2,k3,k4,k5,comment))

# 8. rubric
rub=[('K1','Примитивы','Только слова из списка 65 примитивов; грамматика примитивов верна','1–2 слова вне списка или ошибки в грамматике примитивов','≥3 слов вне списка, толкование не на языке промпта или смесь языков'),
('K2','Ядро значения','Очень хочет того, чего нет, что далеко / может быть недостижимо; длится долго','Только «хочет и нет»','Ядро не схвачено или искажено'),
('K3','Различение','Видно отличие от געגועים и תשוקה','Различие намечено','Подходит к любому «желанию / тоске»'),
('K4','Охват доменов','Общее для религиозной, политической, эротической, психологической сфер; сферы названы','Сползает в одну сферу или пропущены важные','Про одну сферу'),
('K5','Сценарий и формат','Прототипический сценарий Вежбицкой, короткие строки, чистый JSON','Сценарий неполный или строки длинные','Нет сценария или не тот формат')]
cur.executemany('insert into rubric values (?,?,?,?,?)',rub)
con.commit()
for t in ['corpus','prompt_contexts','primes','prompts','models','runs','scores','rubric']:
    print(t, cur.execute(f'select count(*) from {t}').fetchone()[0])
print(cur.execute('select model,lang,n,refusals,k1_mean,k1_sd from summary').fetchall())
con.close()
